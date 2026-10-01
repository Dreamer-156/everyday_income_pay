"""Bill adapters; monetary values use integer cents."""
import csv
import io
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

FIELDS = {
    'time': ['交易时间', '交易日期', '记账日期', '交易发生时间', '入账日期', '记账时间', '日期', '时间'],
    'amount': ['金额(元)', '金额（元）', '交易金额', '交易金额(元)', '交易金额（元）', '发生金额', '金额', '发生额'],
    'direction': ['收/支', '收支', '收支类型', '借贷标志', '收支方向'],
    'merchant': ['交易对方', '对方户名', '交易对象', '商户名称', '商户', '对方名称'],
    'description': ['商品', '商品说明', '交易摘要', '摘要', '备注', '说明'],
    'type': ['交易类型', '业务类型', '交易名称'],
    'status': ['交易状态', '当前状态', '状态'],
    'funding': ['支付方式', '收/付款方式', '付款方式', '资金渠道'],
    'external_id': ['交易单号', '交易订单号', '流水号', '银行流水号', '交易号'],
    'merchant_id': ['商户单号', '商家订单号', '对方流水号'],
    'income_amount': ['收入金额', '贷方发生额', '存入金额', '收入金额(元)', '收入金额（元）', '贷方金额', '收入'],
    'expense_amount': ['支出金额', '借方发生额', '支取金额', '支出金额(元)', '支出金额（元）', '借方金额', '支出'],
}


def clean(value):
    return str(value if value is not None else '').replace('\ufeff', '').strip().strip('\t')


def cents(value):
    value = clean(value).replace(',', '').replace('￥', '').replace('¥', '').replace('元', '')
    if value.startswith('(') and value.endswith(')'):
        value = '-' + value[1:-1]
    try:
        number = Decimal(value)
        if not number.is_finite() or number * 100 != (number * 100).to_integral_value():
            raise ValueError('金额必须精确到分')
        return int(number * 100)
    except InvalidOperation as exc:
        raise ValueError('无法识别金额') from exc


def timestamp(value):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    text = clean(value).replace('/', '-').replace('年', '-').replace('月', '-').replace('日', '').replace('T', ' ')
    for pattern in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d', '%Y%m%d %H%M%S', '%Y%m%d']:
        try:
            return datetime.strptime(text, pattern).strftime('%Y-%m-%d %H:%M:%S')
        except ValueError:
            pass
    raise ValueError('无法识别日期，请使用 YYYY-MM-DD HH:MM:SS')


def guess_mapping(headers):
    mapping = {}
    for field, choices in FIELDS.items():
        for choice in choices:
            if choice in headers:
                mapping[field] = choice
                break
    return mapping


def read_table(name, blob, header_row=None, sheet=None, password=None):
    suffix = Path(name).suffix.lower()
    if suffix == '.pdf':
        from pdf_importer import read_pdf_table
        return read_pdf_table(blob, sheet=sheet, password=password)
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('解压后的表格超过 100 MB，请缩小日期范围')
        workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True, keep_links=False)
        sheets = workbook.sheetnames
        chosen = sheet or sheets[0]
        if chosen not in sheets:
            workbook.close()
            raise ValueError('工作表不存在')
        try:
            rows = []
            for row in workbook[chosen].iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > 100000:
                    raise ValueError('最多支持十万行，请分批导入')
        finally:
            workbook.close()
    elif suffix in {'.csv', '.txt', '.tsv'}:
        text = None
        for encoding in ['utf-8-sig', 'gb18030', 'utf-16']:
            try:
                text = blob.decode(encoding)
                break
            except UnicodeError:
                pass
        if text is None:
            raise ValueError('无法识别文件编码，请另存为 UTF-8 CSV')
        lines = text.splitlines()
        delimiter = '\t' if suffix == '.tsv' else ','
        if suffix != '.tsv' and max((line.count('\t') for line in lines[:80]), default=0) > max((line.count(',') for line in lines[:80]), default=0):
            delimiter = '\t'
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        sheets, chosen = [], ''
        if len(rows) > 100000:
            raise ValueError('最多支持十万行，请分批导入')
    else:
        raise ValueError('支持 CSV、TSV、XLSX、文字表格 PDF；旧 XLS、加密压缩包请先转换或解压')
    if not rows:
        raise ValueError('文件没有内容')
    if header_row is None:
        scored = []
        for i, row in enumerate(rows[:80]):
            mapping = guess_mapping([clean(cell) for cell in row])
            score = len(mapping) + (5 if 'time' in mapping else 0) + (3 if 'amount' in mapping else 0)
            scored.append((score, -i, i))
        index = max(scored)[2]
    else:
        index = int(header_row) - 1
    if not 0 <= index < len(rows):
        raise ValueError('表头行超出文件范围')
    headers = [clean(cell) or f'空列{i+1}' for i, cell in enumerate(rows[index])]
    if len(set(headers)) != len(headers):
        raise ValueError('表头存在重复列名，请先在副本中改名')
    records = []
    for i, row in enumerate(rows[index+1:], index+2):
        if not any(clean(cell) for cell in row):
            continue
        records.append((i, dict(zip(headers, [clean(cell) for cell in row] + [''] * max(0, len(headers)-len(row))))))
    return {'headers': headers, 'mapping': guess_mapping(headers), 'records': records,
            'header_row': index+1, 'sheet': chosen, 'sheets': sheets}


def normalize(row, mapping):
    get = lambda field: clean(row.get(mapping.get(field, ''), ''))
    for column in ['币种', '货币', '币别', 'Currency']:
        currency = clean(row.get(column, '')).upper()
        if currency and currency not in {'CNY', 'RMB', '人民币', '人民币元', '156'}:
            raise ValueError('第一版仅支持人民币，此行币种未识别或为外币')
    when = timestamp(get('time'))
    review = ''
    direction = get('direction')
    if mapping.get('income_amount') and mapping.get('expense_amount'):
        incoming, outgoing = get('income_amount'), get('expense_amount')
        inc = cents(incoming) if incoming and incoming not in {'-', '--'} else 0
        out = cents(outgoing) if outgoing and outgoing not in {'-', '--'} else 0
        if inc and out:
            raise ValueError('同一行收入和支出均非零，请检查列映射')
        amount = abs(inc or out)
        kind = 'income' if inc else 'expense' if out else 'ignored'
    else:
        signed = cents(get('amount'))
        amount = abs(signed)
        if direction in {'收入', '入', '贷', '贷方', '存入', '收', '入账'}:
            kind = 'income'
        elif direction in {'支出', '出', '借', '借方', '支取', '付', '出账'}:
            kind = 'expense'
        elif direction in {'不计收支', '不计收入支出', '/', '中性'}:
            kind = 'unknown'
            review = '渠道标为不计收支，请确认是否属于自己的账户转移'
        elif not direction and signed < 0:
            kind = 'expense'
        else:
            kind = 'unknown'
            review = '收支方向不明，请确认；当前未计入收支'
    desc, tx_type, status = get('description'), get('type'), get('status')
    text = ' '.join([tx_type, desc])
    if any(word in status for word in ['失败', '已关闭', '已取消', '未支付', '待支付']):
        kind, review = 'ignored', ''
    elif ('退款' in tx_type or '退货' in tx_type or (kind == 'income' and '退款' in desc)):
        kind, review = 'refund', ''
    elif '退款' in status and kind == 'expense':
        review = '原付款显示退款状态，请核对独立退款流水；未自动扣减'
    return dict(time=when, cents=amount, direction=direction, base_kind=kind, kind=kind,
                merchant=get('merchant'), description=desc, tx_type=tx_type, status=status,
                funding=get('funding'), external_id='' if get('external_id') in {'/', '-', '--', '0'} else get('external_id'), merchant_id='' if get('merchant_id') in {'/', '-', '--', '0'} else get('merchant_id'),
                category='手续费' if any(w in text for w in ['手续费', '服务费']) else '其他',
                review=review, initial_review=review)
