import base64
import csv
import io
import json
import sqlite3
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import ROOT, Server, export_rows
from core import Ledger, summary
from importer import cents, normalize, read_table

HEAD = ['交易时间','交易类型','交易对方','说明','收支方向','交易金额','支付方式','状态','交易单号','商家订单号']


def bill(*rows):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(HEAD)
    writer.writerows(rows)
    return out.getvalue().encode('utf-8-sig')


def tx(time='2026-09-01 10:00:00', kind='消费', merchant='商户', desc='商品', direction='支出', amount='10.00', funding='微信零钱', status='成功', xid='id-1', order=''):
    return [time,kind,merchant,desc,direction,amount,funding,status,xid,order]


class LedgerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.ledger = Ledger(self.tmp.name)
        self.ledger.account('招商银行(1234)', ['招商银行（1234）'])
    def tearDown(self):
        self.tmp.cleanup()
    def add(self, rows, channel='微信', account=1, name='bill.csv'):
        return self.ledger.import_file(name, bill(*rows), channel, account)
    def rows(self):
        return self.ledger.snapshot()['transactions']
    def total(self, month=''):
        return summary(self.rows(), month)

    def test_repeat_same_file_does_not_add(self):
        result=self.add([tx()])
        again=self.add([tx()])
        self.assertEqual(result['new'],1)
        self.assertTrue(again['repeated_file'])
        self.assertEqual(again['new'],0)
        self.assertEqual(len(self.rows()),1)

    def test_overlapping_files_keep_provenance(self):
        self.add([tx()])
        result=self.add([tx(),tx(xid='id-2',amount='20.00')],name='second.csv')
        self.assertEqual((result['new'],result['duplicate']),(1,1))
        original=next(r for r in self.rows() if r['external_id']=='id-1')
        self.assertEqual(len(original['sources']),2)
        self.assertEqual(self.total()['spending'],3000)

    def test_identical_rows_without_ids_are_not_collapsed_within_file(self):
        self.add([tx(xid=''),tx(xid='')])
        self.assertEqual(len(self.rows()),2)
        self.assertEqual(self.total()['spending'],2000)

    def test_single_side_owned_withdrawal_is_excluded(self):
        self.add([tx(kind='零钱提现',merchant='招商银行(1234)',desc='提现到招商银行(1234)',direction='不计收支',amount='500.00')])
        self.assertEqual(self.rows()[0]['kind'],'transfer')
        self.assertEqual(self.total()['net_spending'],0)
        self.assertFalse(self.rows()[0]['review'])

    def test_unknown_withdrawal_requires_confirmation(self):
        self.add([tx(kind='提现',merchant='银行卡',desc='提现',direction='不计收支')])
        self.assertEqual(self.rows()[0]['kind'],'unknown')
        self.assertTrue(self.rows()[0]['review'])
        self.ledger.edit(self.rows()[0]['id'],{'kind':'transfer'})
        self.assertFalse(self.rows()[0]['review'])

    def test_owned_transfer_crosses_days_and_months(self):
        self.add([tx(time='2026-08-31 23:59:00',kind='提现',merchant='招商银行(1234)',desc='提现',direction='不计收支',amount='500.00')])
        self.add([tx(time='2026-09-01 08:00:00',kind='提现到账',merchant='微信零钱',desc='微信零钱提现',direction='收入',amount='500.00',funding='',xid='bank-1')], '银行卡',3)
        self.assertTrue(all(r['kind']=='transfer' and r['link_id'] for r in self.rows()))
        self.assertEqual(self.total('2026-09')['income'],0)

    def test_bank_to_alipay_transfer(self):
        self.add([tx(kind='余额充值',merchant='招商银行(1234)',desc='充值',direction='不计收支',funding='招商银行(1234)')],'支付宝',2)
        self.add([tx(kind='充值',merchant='支付宝余额',desc='支付宝余额充值',funding='',xid='bank-1')],'银行卡',3)
        self.assertEqual(self.total()['transfers'],2)
        self.assertEqual(self.total()['net_spending'],0)

    def test_same_amount_salary_is_not_transfer_partner(self):
        self.add([tx(kind='提现',merchant='招商银行(1234)',desc='提现',direction='不计收支',amount='500.00')])
        self.add([tx(kind='工资',merchant='单位',desc='工资',direction='收入',amount='500.00',funding='',xid='salary')],'银行卡',3)
        salary=next(r for r in self.rows() if r['external_id']=='salary')
        self.assertEqual(salary['kind'],'income')
        self.assertIsNone(salary['link_id'])

    def test_transfer_fee_and_different_received_amount(self):
        self.add([tx(kind='提现',merchant='招商银行(1234)',desc='提现',direction='不计收支',amount='500.00'),
                  tx(kind='提现手续费',merchant='微信支付',desc='提现手续费',amount='1.00',xid='fee')])
        self.add([tx(kind='提现到账',merchant='微信零钱',desc='微信零钱提现',direction='收入',amount='499.00',funding='',xid='bank-1')],'银行卡',3)
        self.assertEqual(self.total()['spending'],100)
        self.assertEqual(self.total()['income'],0)
        wallet=next(r for r in self.rows() if r['external_id']=='id-1')
        bank=next(r for r in self.rows() if r['external_id']=='bank-1')
        self.ledger.link(wallet['id'],bank['id'],'transfer')
        self.assertEqual(self.total()['spending'],100)

    def test_card_funded_transfer_to_friend_is_real_expense(self):
        self.add([tx(kind='转账',merchant='朋友',desc='饭钱',funding='招商银行(1234)')])
        self.assertEqual(self.rows()[0]['kind'],'expense')
        self.assertEqual(self.total()['spending'],1000)

    def test_phone_topup_is_consumption_not_wallet_recharge(self):
        self.add([tx(kind='商户消费',merchant='运营商',desc='手机话费充值',funding='招商银行(1234)')])
        self.assertEqual(self.rows()[0]['kind'],'expense')
        self.assertEqual(self.total()['spending'],1000)

    def test_foreign_currency_is_reported_not_summed_as_rmb(self):
        data='交易时间,交易金额,收支方向,币种\n2026-09-01,100,支出,USD\n'.encode()
        result=self.ledger.import_file('foreign.csv',data,'银行卡',3)
        self.assertEqual(result['new'],0)
        self.assertEqual(len(result['errors']),1)

    def test_external_person_transfer_remains_income(self):
        self.add([tx(kind='转账',merchant='朋友',desc='转账',direction='收入',amount='100.00')])
        self.assertEqual(self.rows()[0]['kind'],'income')
        self.assertEqual(self.total()['income'],10000)

    def test_duplicate_consumption_with_shared_reference(self):
        self.add([tx(amount='28.00',funding='招商银行(1234)',order='order-1')])
        self.add([tx(time='2026-09-01 10:00:10',kind='快捷支付',merchant='财付通',desc='微信支付',amount='28.00',funding='',xid='bank-1',order='order-1')],'银行卡',3)
        self.assertEqual(self.total()['spending'],2800)
        self.assertEqual(self.total()['duplicates'],1)
        self.assertTrue(all(r['link_id'] for r in self.rows()))

    def test_without_order_reference_needs_manual_confirmation(self):
        self.add([tx(funding='招商银行(1234)')])
        self.add([tx(kind='快捷支付',merchant='财付通',desc='微信支付',funding='',xid='bank-1')],'银行卡',3)
        self.assertEqual(self.total()['spending'],2000)
        self.assertTrue(all(r['review'] for r in self.rows()))
        wallet=next(r for r in self.rows() if r['channel']=='微信')
        bank=next(r for r in self.rows() if r['channel']=='银行卡')
        self.ledger.link(wallet['id'],bank['id'],'duplicate')
        self.assertEqual(self.total()['spending'],1000)
        self.ledger.rule('商品','购物')
        self.assertEqual(self.total()['spending'],1000)
        self.assertFalse(any(r['review'] for r in self.rows()))

    def test_duplicate_refund_is_offset_only_once(self):
        self.add([tx(kind='退款',direction='收入',funding='招商银行(1234)',order='refund-order')])
        self.add([tx(time='2026-09-01 10:00:10',kind='退款',merchant='财付通',desc='微信支付退款',direction='收入',funding='',xid='bank-refund',order='refund-order')],'银行卡',3)
        self.assertEqual(self.total()['refunds'],1000)
        self.assertEqual(self.total()['duplicates'],1)

    def test_refund_category_tracks_purchase_unless_manually_edited(self):
        self.add([tx(),tx(kind='退款',direction='收入',amount='5.00',xid='refund')])
        purchase=next(r for r in self.rows() if r['kind']=='expense')
        refund=next(r for r in self.rows() if r['kind']=='refund')
        self.ledger.link(purchase['id'],refund['id'],'refund')
        self.ledger.edit(purchase['id'],{'category':'学习'})
        self.assertEqual(next(r for r in self.rows() if r['id']==refund['id'])['category'],'学习')
        self.ledger.edit(refund['id'],{'category':'餐饮'})
        self.ledger.edit(purchase['id'],{'category':'购物'})
        self.assertEqual(next(r for r in self.rows() if r['id']==refund['id'])['category'],'餐饮')

    def test_equal_amount_and_time_alone_do_not_deduplicate(self):
        self.add([tx(funding='微信零钱')])
        self.add([tx(merchant='另一个商户',desc='消费',funding='',xid='bank-1')],'银行卡',3)
        self.assertEqual(self.total()['spending'],2000)
        self.assertEqual(self.total()['duplicates'],0)

    def test_ambiguous_transfer_does_not_choose_partner(self):
        self.add([tx(kind='提现',merchant='招商银行(1234)',desc='提现',direction='不计收支',amount='500.00'),
                  tx(time='2026-09-01 10:00:20',kind='提现',merchant='招商银行(1234)',desc='提现',direction='不计收支',amount='500.00',xid='withdraw-2')])
        self.add([tx(kind='提现到账',merchant='微信零钱',desc='微信零钱提现',direction='收入',amount='500.00',funding='',xid='bank-1')],'银行卡',3)
        self.assertTrue(all(r['link_id'] is None for r in self.rows()))
        self.assertTrue(any(r['review'] for r in self.rows()))

    def test_refund_is_not_income_and_uses_arrival_month(self):
        self.add([tx(time='2026-08-31 10:00:00',amount='60.00',status='部分退款'),
                  tx(time='2026-09-01 10:00:00',kind='退款',direction='收入',amount='10.00',xid='refund')])
        self.assertEqual(self.total('2026-08')['spending'],6000)
        self.assertEqual(self.total('2026-09')['net_spending'],-1000)
        self.assertEqual(self.total('2026-09')['income'],0)
        purchase=next(r for r in self.rows() if r['base_kind']=='expense')
        refund=next(r for r in self.rows() if r['base_kind']=='refund')
        self.ledger.edit(purchase['id'],{'category':'学习'})
        self.ledger.link(purchase['id'],refund['id'],'refund')
        self.assertEqual(next(r for r in self.rows() if r['id']==refund['id'])['category'],'学习')

    def test_failed_transactions_are_excluded(self):
        self.add([tx(status='已关闭'),tx(status='支付失败',xid='failed-2')])
        self.assertEqual(self.total()['spending'],0)

    def test_unknown_positive_bank_direction_never_assumed_income(self):
        self.add([tx(direction='',funding='',amount='20.00')],'银行卡',3)
        self.assertEqual(self.total()['income'],0)
        self.assertTrue(self.rows()[0]['review'])

    def test_error_rows_are_reported_and_raw_source_preserved(self):
        data=bill(tx(),tx(amount='NaN',xid='bad'),tx(time='日期错误',xid='bad-date'))
        result=self.ledger.import_file('bill.csv',data,'微信',1)
        self.assertEqual(len(result['errors']),2)
        self.assertEqual(result['new']+result['duplicate']+len(result['errors']),result['rows'])
        self.assertEqual(next(self.ledger.raw_dir.iterdir()).read_bytes(),data)

    def test_rules_never_override_manual_categories(self):
        self.add([tx(merchant='食堂')])
        self.ledger.rule('食堂','餐饮')
        row=self.rows()[0]
        self.assertEqual(row['category'],'餐饮')
        self.ledger.edit(row['id'],{'category':'学习'})
        self.ledger.rule('食堂','购物')
        self.assertEqual(self.rows()[0]['category'],'学习')

    def test_unlink_does_not_immediately_rematch(self):
        self.test_duplicate_consumption_with_shared_reference()
        row=self.rows()[0]
        self.ledger.unlink(row['id'])
        self.ledger.rule('商户','购物')
        self.assertTrue(all(r['link_id'] is None for r in self.rows()))
        self.assertEqual(self.total()['spending'],5600)

    def test_backup_restore_and_invalid_backup_preserve_current(self):
        self.add([tx()])
        backup=self.ledger.backup().read_bytes()
        self.ledger.add_manual(dict(time='2026-09-02',amount='5.00',kind='expense',category='其他',account_id=1))
        self.assertEqual(len(self.rows()),2)
        with self.assertRaises(ValueError):
            self.ledger.restore(b'not a database')
        self.assertEqual(len(self.rows()),2)
        result=self.ledger.restore(backup)
        self.assertEqual(len(self.rows()),1)
        self.assertTrue((self.ledger.root/'backups'/result['previous_backup']).exists())

    def test_money_precision(self):
        self.assertEqual(cents('￥1,234.56'),123456)
        for value in ['1.001','NaN','Infinity','abc']:
            with self.assertRaises(ValueError):
                cents(value)

    def test_gb18030_and_split_income_expense_columns(self):
        text='交易日期,收入金额,支出金额,摘要\n2026-09-01,100,,工资\n2026-09-02,,25,食堂\n'
        result=self.ledger.import_file('bank.csv',text.encode('gb18030'),'银行卡',3)
        self.assertEqual(result['new'],2)
        self.assertEqual((self.total()['income'],self.total()['spending']),(10000,2500))

    def test_xlsx_with_preface_and_native_dates(self):
        from datetime import datetime
        from openpyxl import Workbook
        w=Workbook();s=w.active
        s.append(['银行流水']);s.append(HEAD)
        s.append(tx(time=datetime(2026,9,1,10,0,0)))
        out=io.BytesIO();w.save(out)
        self.ledger.import_file('bill.xlsx',out.getvalue(),'微信',1)
        self.assertEqual(self.rows()[0]['time'],'2026-09-01 10:00:00')

    def test_exports_include_classification_and_escape_formulas(self):
        self.add([tx(merchant='=HYPERLINK(\"bad\")')])
        csv_blob=export_rows(self.rows())
        self.assertIn("'=HYPERLINK",csv_blob.decode('utf-8-sig'))
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(export_rows(self.rows(),xlsx=True)))
        self.assertNotEqual(workbook.active.cell(2,8).data_type,'f')


class HttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory()
        cls.server=Server(('127.0.0.1',0),cls.tmp.name)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True)
        cls.thread.start()
        cls.url=f'http://127.0.0.1:{cls.server.server_port}'
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();cls.thread.join();cls.tmp.cleanup()
    def request(self,path,body=None,token=True,origin=None):
        headers={}
        if token: headers['X-Ledger-Token']=self.server.token
        if origin: headers['Origin']=origin
        payload=json.dumps(body).encode() if body is not None else None
        req=urllib.request.Request(self.url+path,data=payload,headers=headers)
        return urllib.request.urlopen(req,timeout=15)
    def test_health_is_private_data_free_and_available_without_token(self):
        with self.request('/health',token=False) as response:
            health=json.load(response)
        self.assertEqual(health['app'],'local-ledger')
        self.assertEqual(set(health),{'app','instance','port','pid'})
        from launch import healthy
        self.assertTrue(healthy(self.url))
        self.assertFalse(healthy(self.url,instance='another-project'))
        with self.assertRaises(urllib.error.HTTPError):
            self.request('/health',token=False,origin='https://example.com')

    def test_strict_port_does_not_silently_change_address(self):
        import subprocess
        import os
        result=subprocess.run([sys.executable,str(ROOT/'app.py'),'--port',str(self.server.server_port),
                               '--strict-port','--no-browser','--data-dir',str(Path(self.tmp.name)/'port-test')],
                              capture_output=True,timeout=15,env=dict(os.environ,PYTHONIOENCODING='utf-8'))
        self.assertNotEqual(result.returncode,0)
        self.assertIn('端口不可用',result.stderr.decode('utf-8',errors='replace'))

    def test_requires_token_and_rejects_cross_origin(self):
        for kwargs in [{'token':False},{'origin':'https://example.com'}]:
            with self.assertRaises(urllib.error.HTTPError) as cm:
                self.request('/api/state',**kwargs)
            self.assertEqual(cm.exception.code,403)
    def test_page_and_demo_are_separate(self):
        with self.request('/') as response:
            page=response.read().decode()
        self.assertIn('收支簿',page)
        self.assertNotIn('__TOKEN__',page)
        with self.request('/demo/api/state') as response:
            demo=json.load(response)
        stats=summary(demo['transactions'])
        self.assertEqual((stats['income'],stats['net_spending'],stats['balance']),(300000,7900,292100))
        self.assertEqual(stats['duplicates'],1)
        self.assertEqual(stats['transfers'],6)
        self.assertEqual(stats['pending'],0)
        with self.request('/api/state') as response:
            self.assertEqual(json.load(response)['transactions'],[])
    def test_preview_import_repeat_and_export_http(self):
        content=base64.b64encode(bill(tx())).decode()
        params={'name':'bill.csv','content':content}
        with self.request('/api/preview',params) as response:
            preview=json.load(response)
        self.assertEqual(preview['count'],1)
        params.update(channel='微信',account_id=1,mapping=preview['mapping'])
        with self.request('/api/import',params) as response:
            self.assertEqual(json.load(response)['new'],1)
        with self.request('/api/import',params) as response:
            self.assertTrue(json.load(response)['repeated_file'])
        with self.request('/api/export.xlsx') as response:
            self.assertTrue(response.read().startswith(b'PK'))


if __name__=='__main__':
    unittest.main()
